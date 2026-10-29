# main.py
import datetime
import io
import os
from typing import Optional, List
import requests
from fastapi import FastAPI, Request, Response, HTTPException, Depends
from fastapi.responses import PlainTextResponse, StreamingResponse
from pydantic import BaseModel
from pydantic_settings import BaseSettings, SettingsConfigDict
from sqlalchemy import create_engine, Column, Integer, String, Float, DateTime, ForeignKey, Boolean
from sqlalchemy.ext.declarative import declarative_base
from sqlalchemy.orm import sessionmaker, Session
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

# --- CONFIGURACIÓN DE SEGURIDAD (VARIABLES DE ENTORNO) ---
class Settings(BaseSettings):
    database_url: str = "sqlite:///subasta_agricola.db"
    webhook_verify_token: str = "Alimentos_HCHF4_Secret_Token_2026"
    whatsapp_access_token: str = "TOKEN_POR_DEFECTO"
    phone_number_id: str = "ID_POR_DEFECTO"
    
    @property
    def meta_api_url(self) -> str:
        return f"https://graph.facebook.com/v18.0/{self.phone_number_id}/messages"
    
    model_config = SettingsConfigDict(env_file=".env", env_file_encoding="utf-8")

settings = Settings()

# --- ARQUITECTURA DE BASE DE DATOS (SQLAlchemy) ---
engine = create_engine(settings.database_url, connect_args={"check_same_thread": False} if "sqlite" in settings.database_url else {})
SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
Base = declarative_base()

class EmpresaAdministradora(Base):
    __tablename__ = "empresa_administradora"
    id = Column(Integer, primary_key=True)
    nombre = Column(String, default="Alimentos HCHF4 C.A.")
    email = Column(String, default="alimentoshchf4@gmail.com")
    telefono = Column(String, default="04120582064")

class UsuarioEstado(Base):
    __tablename__ = "usuario_estados"
    telefono = Column(String, primary_key=True)
    legal_aceptado = Column(Boolean, default=False)
    estado_flujo = Column(String, default="INICIO")  # INICIO, ESPERANDO_INSUMO, ESPERANDO_PRECIO
    municipio_seleccionado = Column(String, nullable=True)
    subasta_activa_id = Column(Integer, nullable=True)

class Agrotienda(Base):
    __tablename__ = "agrotiendas"
    id = Column(Integer, primary_key=True, index=True)
    nombre = Column(String, nullable=False)
    telefono = Column(String, unique=True, nullable=False)
    municipio = Column(String, nullable=False)             
    categoria = Column(String, nullable=False)             
    creditos = Column(Integer, default=10)

class SubastaUrgente(Base):
    __tablename__ = "subastas_urgentes"
    id = Column(Integer, primary_key=True, index=True)
    agricultor_telefono = Column(String, nullable=False)
    municipio = Column(String, nullable=False)
    insumo_requerido = Column(String, nullable=False)
    estado = Column(String, default="ABIERTA")             
    fecha_creacion = Column(DateTime, default=datetime.datetime.utcnow)

class OfertaTienda(Base):
    __tablename__ = "ofertas_tiendas"
    id = Column(Integer, primary_key=True, index=True)
    subasta_id = Column(Integer, ForeignKey("subastas_urgentes.id"))
    agrotienda_id = Column(Integer, ForeignKey("agrotiendas.id"))
    precio_ofrecido = Column(Float, nullable=False)
    fecha_oferta = Column(DateTime, default=datetime.datetime.utcnow)

Base.metadata.create_all(bind=engine)

def get_db():
    db = SessionLocal()
    try: yield db
    finally: db.close()

# --- INICIALIZACIÓN DE FASTAPI ---
app = FastAPI(title="Alimentos HCHF4 C.A. - Ecosistema Yaracuy", description="Sistema de Subasta Inversa Agrícola")

@app.on_event("startup")
def inicializar_sistema():
    db = SessionLocal()
    if not db.query(EmpresaAdministradora).filter_by(email="alimentoshchf4@gmail.com").first():
        db.add(EmpresaAdministradora(nombre="Alimentos HCHF4 C.A.", email="alimentoshchf4@gmail.com", telefono="04120582064"))
        db.commit()
    db.close()

# --- FUNCIONES AUXILIARES WHATSAPP ---
def enviar_mensaje_whatsapp(telefono_destino: str, texto: str):
    headers = {"Authorization": f"Bearer {settings.whatsapp_access_token}", "Content-Type": "application/json"}
    data = {"messaging_product": "whatsapp", "recipient_type": "individual", "to": telefono_destino, "type": "text", "text": {"preview_url": False, "body": texto}}
    try:
        requests.post(settings.meta_api_url, headers=headers, json=data)
    except Exception as e:
        print(f"Error enviando WhatsApp: {e}")

def enviar_consentimiento_legal(telefono_destino: str):
    headers = {"Authorization": f"Bearer {settings.whatsapp_access_token}", "Content-Type": "application/json"}
    data = {
        "messaging_product": "whatsapp", "recipient_type": "individual", "to": telefono_destino, "type": "interactive",
        "interactive": {
            "type": "button",
            "body": {"text": "⚖️ *Aviso Legal y Términos* \n\nPara operar en la red de subastas de Alimentos HCHF4 C.A., debes aceptar nuestros términos y condiciones."},
            "action": {"buttons": [{"type": "reply", "reply": {"id": "legal_aceptado", "title": "Sí, Acepto"}}, {"type": "reply", "reply": {"id": "legal_rechazado", "title": "No Acepto"}}]}
        }
    }
    try: requests.post(settings.meta_api_url, headers=headers, json=data)
    except Exception as e: print(f"Error: {e}")

def enviar_botones_municipios(telefono_destino: str):
    headers = {"Authorization": f"Bearer {settings.whatsapp_access_token}", "Content-Type": "application/json"}
    data = {
        "messaging_product": "whatsapp", "recipient_type": "individual", "to": telefono_destino, "type": "interactive",
        "interactive": {
            "type": "button",
            "body": {"text": "🌾 Bienvenido a Alimentos HCHF4. Selecciona tu municipio en Yaracuy:"},
            "action": {"buttons": [{"type": "reply", "reply": {"id": "mun_san_felipe", "title": "San Felipe"}}, {"type": "reply", "reply": {"id": "mun_chivacoa", "title": "Bruzual (Chivacoa)"}}, {"type": "reply", "reply": {"id": "mun_yaritagua", "title": "Peña (Yaritagua)"}}]}
        }
    }
    try: requests.post(settings.meta_api_url, headers=headers, json=data)
    except Exception as e: print(f"Error: {e}")

def enviar_alerta_subasta_a_tienda(telefono_tienda: str, subasta_id: int, insumo: str):
    headers = {"Authorization": f"Bearer {settings.whatsapp_access_token}", "Content-Type": "application/json"}
    data = {
        "messaging_product": "whatsapp", "recipient_type": "individual", "to": telefono_tienda, "type": "interactive",
        "interactive": {
            "type": "button",
            "header": {"type": "text", "text": "🚨 ALERTA DE SUBASTA URGENTE"},
            "body": {"text": f"Un productor solicita:\n📦 *Insumo:* {insumo}\n\n¿Deseas ofertar? (Descuenta 1 crédito)."},
            "action": {"buttons": [{"type": "reply", "reply": {"id": f"sub_aceptar_{subasta_id}", "title": "Sí, Ofertar"}}, {"type": "reply", "reply": {"id": f"sub_rechazar_{subasta_id}", "title": "No participar"}}]}
        }
    }
    try: requests.post(settings.meta_api_url, headers=headers, json=data)
    except Exception as e: print(f"Error: {e}")

# --- WEBHOOKS OFICIALES META API ---
@app.get("/webhook/whatsapp", response_class=PlainTextResponse, tags=["[WEBHOOK] Meta"])
async def verificar_webhook(request: Request):
    params = request.query_params
    if params.get("hub.mode") == "subscribe" and params.get("hub.verify_token") == settings.webhook_verify_token:
        return params.get("hub.challenge")
    raise HTTPException(status_code=403, detail="Verification token mismatch")

@app.post("/webhook/whatsapp", tags=["[WEBHOOK] Meta"])
async def recibir_mensaje_whatsapp(request: Request, db: Session = Depends(get_db)):
    payload = await request.json()
    try:
        if "object" in payload and payload["object"] == "whatsapp_business_account":
            for entry in payload.get("entry", []):
                for change in entry.get("changes", []):
                    value = change.get("value", {})
                    if "messages" in value:
                        for message in value["messages"]:
                            telefono = message["from"]
                            tipo = message["type"]
                            
                            user_state = db.query(UsuarioEstado).filter_by(telefono=telefono).first()
                            if not user_state:
                                user_state = UsuarioEstado(telefono=telefono)
                                db.add(user_state); db.commit()
                            
                            if tipo == "interactive" and message["interactive"]["type"] == "button_reply":
                                id_boton = message["interactive"]["button_reply"]["id"]
                                titulo_boton = message["interactive"]["button_reply"]["title"]
                                
                                if id_boton == "legal_aceptado":
                                    user_state.legal_aceptado = True; db.commit()
                                    enviar_mensaje_whatsapp(telefono, "¡Gracias! Cuenta habilitada.")
                                    enviar_botones_municipios(telefono)
                                elif id_boton == "legal_rechazado":
                                    enviar_mensaje_whatsapp(telefono, "Es obligatorio aceptar para operar. Escribe *INICIO* para reintentar.")
                                elif id_boton.startswith("mun_"):
                                    user_state.municipio_seleccionado = titulo_boton
                                    user_state.estado_flujo = "ESPERANDO_INSUMO"; db.commit()
                                    enviar_mensaje_whatsapp(telefono, f"Fijado en: {titulo_boton}. Escribe el insumo urgente:")
                                elif id_boton.startswith("sub_aceptar_"):
                                    sub_id = int(id_boton.replace("sub_aceptar_", ""))
                                    tienda = db.query(Agrotienda).filter_by(telefono=telefono).first()
                                    if not tienda or tienda.creditos <= 0:
                                        enviar_mensaje_whatsapp(telefono, "⚠️ No tienes créditos disponibles. Comunícate con Alimentos HCHF4 para recargar.")
                                    else:
                                        user_state.estado_flujo = "ESPERANDO_PRECIO"
                                        user_state.subasta_activa_id = sub_id; db.commit()
                                        enviar_mensaje_whatsapp(telefono, "Responde ÚNICAMENTE con el precio en dólares (Ej: 45.00):")
                                elif id_boton.startswith("sub_rechazar_"):
                                    enviar_mensaje_whatsapp(telefono, "Alerta archivada sin costo.")
                            
                            elif tipo == "text":
                                texto = message["text"]["body"].strip()
                                if texto.lower() in ["hola", "buenas", "inicio"]:
                                    user_state.estado_flujo = "INICIO"; db.commit()
                                    if not user_state.legal_aceptado: enviar_consentimiento_legal(telefono)
                                    else: enviar_botones_municipios(telefono)
                                    continue
                                
                                if user_state.estado_flujo == "ESPERANDO_INSUMO":
                                    nueva_subasta = SubastaUrgente(agricultor_telefono=telefono, municipio=user_state.municipio_seleccionado, insumo_requerido=texto)
                                    db.add(nueva_subasta); user_state.estado_flujo = "INICIO"; db.commit()
                                    enviar_mensaje_whatsapp(telefono, f"🚨 Urgencia enviada a los comercios de {user_state.municipio_seleccionado}.")
                                    tiendas = db.query(Agrotienda).filter_by(municipio=user_state.municipio_seleccionado).all()
                                    for t in tiendas: enviar_alerta_subasta_a_tienda(t.telefono, nueva_subasta.id, texto)
                                
                                elif user_state.estado_flujo == "ESPERANDO_PRECIO":
                                    try:
                                        precio_float = float(texto)
                                        tienda = db.query(Agrotienda).filter_by(telefono=telefono).first()
                                        nueva_oferta = OfertaTienda(subasta_id=user_state.subasta_activa_id, agrotienda_id=tienda.id, precio_ofrecido=precio_float)
                                        tienda.creditos -= 1; user_state.estado_flujo = "INICIO"; db.add(nueva_oferta); db.commit()
                                        enviar_mensaje_whatsapp(telefono, f"✅ Oferta enviada. Créditos restantes: *{tienda.creditos}*.")
                                    except ValueError:
                                        enviar_mensaje_whatsapp(telefono, "Por favor, ingresa un número válido (Ej: 45.00).")
    except Exception as e:
        print(f"Error en Webhook: {e}")
    return Response(status_code=200)

# --- MÓDULO ADMINISTRATIVO Y EXCEL ---
class CrearTiendaSchema(BaseModel):
    nombre: str; telefono: str; municipio: str; categoria: str

class RecargaCreditosSchema(BaseModel):
    telefono_tienda: str; cantidad_creditos: int; referencia_pago_movil: str

@app.post("/admin/agrotiendas", tags=["[ADMINISTRACIÓN] Comercios"])
def admin_registrar_agrotienda(tienda: CrearTiendaSchema, db: Session = Depends(get_db)):
    if db.query(Agrotienda).filter_by(telefono=tienda.telefono).first():
        raise HTTPException(status_code=400, detail="La agrotienda ya está registrada.")
    nueva_tienda = Agrotienda(nombre=tienda.nombre, telefono=tienda.telefono, municipio=tienda.municipio.strip().title(), categoria=tienda.categoria.strip().title(), creditos=10)
    db.add(nueva_tienda); db.commit()
    return {"status": "Éxito", "mensaje": f"Agrotienda '{nueva_tienda.nombre}' registrada con 10 créditos."}

@app.post("/admin/agrotiendas/recargar", tags=["[ADMINISTRACIÓN] Comercios"])
def admin_recargar_creditos(recarga: RecargaCreditosSchema, db: Session = Depends(get_db)):
    tienda = db.query(Agrotienda).filter_by(telefono=recarga.telefono_tienda).first()
    if not tienda: raise HTTPException(status_code=404, detail="Agrotienda no encontrada.")
    tienda.creditos += recarga.cantidad_creditos; db.commit()
    return {"status": "Recarga Exitosa", "agrotienda": tienda.nombre, "creditos_totales": tienda.creditos}

@app.get("/admin/reportes/excel", tags=["[ADMINISTRACIÓN] Reportes"])
def admin_descargar_reporte_excel(db: Session = Depends(get_db)):
    wb = Workbook()
    ws = wb.active
    ws.title = "Historial de Subastas"
    ws.views.sheetView[0].showGridLines = True

    fuente_titulo = Font(name="Arial", size=14, bold=True, color="FFFFFF")
    fuente_cabecera = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    fuente_datos = Font(name="Arial", size=10)
    relleno_titulo = PatternFill(start_color="1B4D3E", end_color="1B4D3E", fill_type="solid")
    relleno_cabecera = PatternFill(start_color="2C5E43", end_color="2C5E43", fill_type="solid")
    alineacion_centro = Alignment(horizontal="center", vertical="center")
    alineacion_izquierda = Alignment(horizontal="left", vertical="center")

    ws.merge_cells("A1:F1")
    ws["A1"] = "REPORTE GENERAL DE SUBASTAS - ALIMENTOS HCHF4 C.A."
    ws["A1"].font = fuente_titulo; ws["A1"].fill = relleno_titulo; ws["A1"].alignment = alineacion_centro
    ws.row_dimensions[1].height = 40

    cabeceras = ["ID Subasta", "Municipio", "Insumo Urgente", "Estado", "Ofertas Recibidas", "Fecha de Creación"]
    ws.append([]); ws.append(cabeceras)
    ws.row_dimensions[3].height = 25
    for col_num, cabecera in enumerate(cabeceras, 1):
        celda = ws.cell(row=3, column=col_num)
        celda.font = fuente_cabecera; celda.fill = relleno_cabecera; celda.alignment = alineacion_centro

    subastas = db.query(SubastaUrgente).all()
    for fila_num, s in enumerate(subastas, 4):
        total_ofertas = db.query(OfertaTienda).filter_by(subasta_id=s.id).count()
        fecha_str = s.fecha_creacion.strftime("%Y-%m-%d %H:%M:%S")
        ws.append([s.id, s.municipio, s.insumo_requerido, s.estado, total_ofertas, fecha_str])
        ws.row_dimensions[fila_num].height = 20
        for col_num in range(1, 7):
            celda = ws.cell(row=fila_num, column=col_num)
            celda.font = fuente_datos
            celda.alignment = alineacion_centro if col_num in [1, 4, 5] else alineacion_izquierda

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    nombre_archivo = f"Reporte_HCHF4_Yaracuy_{datetime.datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={nombre_archivo}"}
    )